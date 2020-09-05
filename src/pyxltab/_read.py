"""
Read the files to be modified.
"""

# TODO: Migrate useful content into `pyxltab`, remove the rest.

__all__ = ["read"]

from glob import glob
import os
from typing import Dict, List

import openpyxl

IN_DIR = os.getcwd()
NUM_OUT = 1
OUT_SUFFIX_BEGIN = 1001


# * ------------------------------------------------------------------------------ # *
# * MAIN * #


def read(
    in_dir: str = IN_DIR,
    num_out: int = NUM_OUT,
    out_suffix_begin: int = OUT_SUFFIX_BEGIN,
) -> Dict:
    """
    Determine the structure of the directory and workbooks contained within.
    """

    file_structure: Dict = {}
    workbook_filenames = glob(os.path.join(in_dir, "[!~$]*.xlsx"))
    workbook_filenames = [os.path.abspath(filename) for filename in workbook_filenames]

    file_structure["books"] = [
        {"in": filename, "out": generate_out_files(filename, num_out, out_suffix_begin)}
        for filename in workbook_filenames
    ]

    # For each desired input book...
    for book_id in file_structure["books"]:
        book = openpyxl.load_workbook(book_id["in"])

        # Get the sheet names
        book_id["sheets"] = [{"name": sheet_name} for sheet_name in book.sheetnames]

        # For each sheet in the book...
        for sheet_id in book_id["sheets"]:
            sheet = book[sheet_id["name"]]

            # Get the table names
            sheet_id["tables"] = [
                {"name": table_name} for table_name in sheet.tables.keys()
            ]

            # For each table in the sheet...
            for table_id in sheet_id["tables"]:
                table = sheet.tables[table_id["name"]]

                # Get the column names
                table_id["columns"] = [column.name for column in table.tableColumns]

    return file_structure


# * ------------------------------------------------------------------------------ # *
# * FUNCTIONS * #


def generate_out_files(filename: str, num_files: int, suffix_begin) -> List[str]:
    """
    Generates output filenames given an input filename.
    """
    (root, ext) = os.path.splitext(filename)
    (folder, file) = os.path.split(root)
    root = os.path.join(folder, file, file)
    out_files = [
        root + "_" + str(suffix_begin + num + 1) + ext for num in range(num_files)
    ]

    return out_files
